"""The parser against an anonymised slice of the real blok 12 tab (fixtures/blok12_slice.json).

Synthetic fixtures only held what we expected; the real sheet holds typed kg as TEXT, formula kg
as numbers, ranges ("8 - 12"), typos ("RPE 7 - 5") and blank weeks. Refresh the slice when the
coach starts a new block (see the note in the JSON).
"""

import datetime as dt
import io
import json
import logging
from collections import Counter
from collections.abc import Mapping
from pathlib import Path

import openpyxl
import pytest

from conftest import SYNTHETIC_SHEET_ID, InMemoryPostgrest
from training_load.cli import collect_strength
from training_load.config import GoogleSettings
from training_load.sources.strength_sheet import (
    KG_NOT_A_NUMBER,
    ROW_WITHOUT_TYPE,
    SETS_REPS_NOT_A_NUMBER,
    ParsedSet,
    parse_all,
)

SLICE = json.loads((Path(__file__).parent / "fixtures" / "blok12_slice.json").read_text("utf-8"))
TAB: str = SLICE["tab"]
BODYWEIGHT = 80.0


def workbook(edits: Mapping[tuple[int, int], object] | None = None) -> bytes:
    """The slice as xlsx, with the real cell types; ``edits`` overrides (row, 0-based col)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = TAB
    values: dict[tuple[int, int], object] = {}
    for row, col, kind, value in SLICE["cells"]:
        values[(row, col)] = dt.datetime.fromisoformat(value) if kind == "d" else value
    values.update(edits or {})
    for (row, col), value in values.items():
        ws.cell(row=row, column=col + 1, value=value)  # type: ignore[call-overload]  # slice cells are str/float/datetime
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def sets_of(parsed: list[ParsedSet], row: int, week: int) -> list[ParsedSet]:
    return [p for p in parsed if p["row"] == row and p["week"] == week]


def test_real_slice_parses_without_issues() -> None:
    issues: Counter[str] = Counter()
    parsed = parse_all(workbook(), BODYWEIGHT, issues)
    assert issues == Counter()
    assert {p["section"] for p in parsed} == {1}
    # Typed kg arrive as text ("110"), formula kg (the -10% row) as a number.
    assert {p["logged_kg"] for p in sets_of(parsed, 25, 1)} == {110.0}
    assert {p["logged_kg"] for p in sets_of(parsed, 27, 1)} == {100.0}
    # Week 2 is prescribed but not logged yet: missing is None, never 0.
    assert {p["logged_kg"] for p in sets_of(parsed, 25, 2)} == {None}
    # The -10% row's formula gives a numeric 0 until the top set is logged: that is missing too.
    assert {p["logged_kg"] for p in sets_of(parsed, 27, 2)} == {None}
    # "8 - 12" and the coach's "RPE 7 - 5" typo still parse.
    assert sets_of(parsed, 31, 1)[0]["reps"] == 10.0
    assert sets_of(parsed, 31, 4)[0]["prescribed"] == "RPE 7 - 5"


def test_unreadable_kg_is_counted_and_stored_as_none() -> None:
    issues: Counter[str] = Counter()
    parsed = parse_all(workbook({(25, 7): "BW"}), BODYWEIGHT, issues)
    assert issues == Counter({KG_NOT_A_NUMBER: 1})
    assert {p["logged_kg"] for p in sets_of(parsed, 25, 1)} == {None}
    # The score formula is unchanged: an unreadable kg scores like a blank one.
    blank = parse_all(workbook({(25, 7): None}), BODYWEIGHT)
    assert [p["score"] for p in sets_of(parsed, 25, 1)] == [
        p["score"] for p in sets_of(blank, 25, 1)
    ]


def test_junk_sets_and_untyped_rows_are_counted_not_crashing_or_vanishing() -> None:
    issues: Counter[str] = Counter()
    parsed = parse_all(workbook({(26, 4): "x", (28, 1): None}), BODYWEIGHT, issues)
    assert issues == Counter({SETS_REPS_NOT_A_NUMBER: 1, ROW_WITHOUT_TYPE: 1})
    assert sets_of(parsed, 26, 1) == []  # skipped, and counted
    assert sets_of(parsed, 28, 1) == []


GOOGLE = GoogleSettings(sheet_id=SYNTHETIC_SHEET_ID, service_account_info={})


def run_on(xlsx: bytes, db: InMemoryPostgrest, monkeypatch: pytest.MonkeyPatch) -> int:
    monkeypatch.setattr(collect_strength, "access_token", lambda _info: "token")
    monkeypatch.setattr(collect_strength, "download_workbook", lambda *_a, **_k: xlsx)
    send = lambda *_a, **_k: None  # noqa: E731  # never called: download_workbook is replaced
    return collect_strength.run(GOOGLE, bodyweight=BODYWEIGHT, send=send, db=db)  # type: ignore[arg-type]


def test_a_few_issues_are_logged_as_counts_and_the_run_writes(
    db: InMemoryPostgrest, monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    with caplog.at_level(logging.WARNING):
        written = run_on(workbook({(25, 7): "BW"}), db, monkeypatch)
    assert written == len(db.tables["strength_sets"]) > 0
    assert "1 cells unreadable: kg not a number" in caplog.text
    assert "BW" not in caplog.text  # counts only: the Actions logs are public
    row = next(r for r in db.tables["strength_sets"] if r["sheet_row"] == 25 and r["week"] == 1)
    assert row["logged_kg"] is None


def test_too_many_issues_fail_before_anything_is_written(
    db: InMemoryPostgrest, monkeypatch: pytest.MonkeyPatch
) -> None:
    run_on(workbook(), db, monkeypatch)
    before = list(db.tables["strength_sets"])
    junk = {(row, 4 + 8 * w): "x" for row in range(25, 35) for w in range(4)}
    with pytest.raises(collect_strength.TooManyIssuesError, match="nothing written"):
        run_on(workbook(junk), db, monkeypatch)
    assert db.tables["strength_sets"] == before
