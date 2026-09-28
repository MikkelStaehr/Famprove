"""Shared fixtures. All data is synthetic: the repo is public, never commit the real workbook.

Synthetic workbook layout (mimics the "Powerlifting Now" template as strength_sheet reads it;
columns are 0-based tuple indices of ``ws.iter_rows(values_only=True)``):

  Tabs: "Program - blok 11 test", "Program - blok 12 test", and "Oversigt" (must be ignored).

  e1RM table, within the first 15 rows and BEFORE the first date row:
      col i = "SQUAT" / "BENCH" / "DEADLIFT", col i + 5 = e1RM kg (number).

  Date row (one per training-day section; a tab has >= 2 sections so a week spans dates):
      col 1 = any datetime, and for week w (0-based): col 4 + 8w = "WEEK {w+1}",
      col 6 + 8w = that week's session datetime. "WEEK 1" must be an exact cell value.

  Header rows to be skipped: col 1 in ("TYPE", "DAY", "LIFT", "MUSCLE GROUP", "Micro Length")
      or starting with "MICRO", or col 2 == "NAME".

  Exercise row: col 1 = type ("SQUAT", "BENCH", "BACK", "QUADS", ...), col 2 = name,
      per week w: col 4 + 8w = sets, 5 + 8w = reps (5 or "8 - 12"),
      6 + 8w = load text ("RPE 7 - 8" or "-10%"), 7 + 8w = logged kg (number or empty).
  ABS row: col 1 empty, col 2 starting with "Abs".
  Bodyweight row: name containing "dips" / "chin" / "pull-up" / "push-up".

  Covers: a filled week, an unfilled (future) week, a week whose ONLY logged kg is on a
  bodyweight or ABS row (not filled), a tempo variant, blok 12 with one filled week (makes
  blok 11 finished), and a date after the test's fixed ``today``.
"""

import io
import json
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass, field, replace
from datetime import date, datetime
from typing import Any

import openpyxl
import pytest
from openpyxl.worksheet.worksheet import Worksheet

from training_load.db.client import Filters, JsonRow, Postgrest
from training_load.domain.strength import StrengthSet
from training_load.http import DEFAULT_TIMEOUT_S, HttpResponse, QueryParams

SYNTHETIC_BODYWEIGHT = 80.0
SYNTHETIC_SHEET_ID = "synthetic-sheet"
FIXED_TODAY = date(2026, 3, 15)

BLOK_11 = "Program - blok 11 test"
BLOK_12 = "Program - blok 12 test"

type WeekCell = tuple[object, object, object, object] | None
"""(sets, reps, load text, logged kg) for one week, or None when nothing is prescribed."""
type Exercise = tuple[str | None, str, Sequence[WeekCell]]
type Section = tuple[Sequence[date], Sequence[Exercise]]

# Sheet rows (1-based) are predictable: 3 e1RM rows, a blank row, then per section
# DAY, date row, TYPE header, MICRO row, exercises.
BLOK_11_SECTIONS: list[Section] = [
    (
        [date(2026, 2, 2), date(2026, 2, 9), date(2026, 2, 16)],
        [
            (
                "SQUAT",
                "Squat",
                [(3, 5, "RPE 7 - 8", 120), (3, 5, "RPE 8", None), (2, 3, "RPE 6", None)],
            ),
            (
                "BACK",
                "Dips",
                [(3, "8 - 12", "RPE 7", 10), (3, 10, "RPE 7", 10), (3, 10, "RPE 7", None)],
            ),
            (None, "Abs rollout", [(2, 10, None, None), (2, 10, None, 5), (2, 10, None, None)]),
            ("BACK", "Chin-ups", [(3, 5, "BW", None), None, None]),
        ],
    ),
    (
        [date(2026, 2, 5), date(2026, 2, 12), date(2026, 2, 19)],
        [
            (
                "BENCH",
                "Tempo bench",
                [(3, 6, "RPE 6", 70), (3, 6, "RPE 7", None), (3, 6, "RPE 7", None)],
            ),
            ("QUADS", "Leg extension", [(3, 12, "RPE 8", 40), (3, 12, "RPE 8", None), None]),
            ("BENCH", "Bench press", [(4, 5, "-10%", None), (4, 5, "-10%", None), None]),
        ],
    ),
]
BLOK_12_SECTIONS: list[Section] = [
    (
        [date(2026, 3, 9), date(2026, 3, 16)],
        [("SQUAT", "Squat", [(3, 5, "RPE 7", 125), (3, 5, "RPE 8", 130)])],
    ),
]


def _write_tab(
    ws: Worksheet, e1rm: Mapping[str, float], sections: Sequence[Section], n_weeks: int
) -> None:
    width = 4 + 8 * n_weeks
    for lift, kg in e1rm.items():
        ws.append([None, lift, None, None, None, None, kg])
    ws.append([])
    for day_no, (dates, exercises) in enumerate(sections, start=1):
        ws.append([None, "DAY", f"Day {day_no}"])
        date_row: list[object] = [None] * width
        date_row[1] = datetime.combine(dates[0], datetime.min.time())
        for w, d in enumerate(dates):
            date_row[4 + 8 * w] = f"WEEK {w + 1}"
            date_row[6 + 8 * w] = datetime.combine(d, datetime.min.time())
        ws.append(date_row)
        # Full template width: the parser indexes every week's 8 columns, even when empty.
        header: list[object] = [None, "TYPE", "NAME", None, "SETS", "REPS", "LOAD", "KG"]
        ws.append([*header, *([None] * (width - 1 - len(header))), "NOTES"])
        ws.append([None, "MICRO 1", "Accumulation"])
        for typ, name, weeks in exercises:
            row: list[object] = [None, typ, name, None]
            for cell in weeks:
                row.extend([*(cell or (None, None, None, None)), None, None, None, None])
            ws.append(row)


def build_workbook() -> bytes:
    wb = openpyxl.Workbook()
    overview = wb.active
    assert overview is not None
    overview.title = "Oversigt"
    overview.append([None, "SQUAT", "Squat", None, 5, 5, "RPE 9", 200])
    _write_tab(
        wb.create_sheet(BLOK_11),
        {"SQUAT": 150.0, "BENCH": 100.0, "DEADLIFT": 180.0},
        BLOK_11_SECTIONS,
        n_weeks=3,
    )
    _write_tab(wb.create_sheet(BLOK_12), {"SQUAT": 160.0}, BLOK_12_SECTIONS, n_weeks=2)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def workbook_bytes() -> bytes:
    """The synthetic workbook above, built in memory with openpyxl and saved to bytes."""
    return build_workbook()


_DEFAULT_SET = StrengthSet(
    sheet_id=SYNTHETIC_SHEET_ID,
    block=BLOK_11,
    sheet_row=10,
    week=1,
    set_no=1,
    date=date(2026, 2, 2),
    type="SQUAT",
    name="Squat",
    reps=5.0,
    logged_kg=100.0,
    kg=100.0,
    bodyweight=False,
    rpe=7.0,
    score=100.0,
)


@pytest.fixture
def make_set() -> Callable[..., StrengthSet]:
    """Factory: StrengthSet with sensible defaults; keyword overrides per test."""

    # Any: overrides span every StrengthSet field type; dataclasses.replace checks them.
    def factory(**overrides: Any) -> StrengthSet:
        return replace(_DEFAULT_SET, **overrides)

    return factory


# --- intervals.icu payloads -------------------------------------------------------------------


def raw_ride(
    activity_id: str, start: str, load: int | None = 80, kind: str = "VirtualRide"
) -> dict[str, object]:
    return {
        "id": activity_id,
        "start_date_local": start,
        "type": kind,
        "name": "Zwift",
        "icu_training_load": load,
        "icu_weighted_avg_watts": 250,
        "icu_intensity": 85.3,
        "icu_ftp": 290,
        "moving_time": 3600,
        "elapsed_time": 3700,
        "power_load": load,
        "hr_load": 70,
        "device_name": "HAMMERHEAD Karoo",
    }


STRAVA_STUB = {
    "id": "12345678",
    "icu_athlete_id": "i123",
    "start_date_local": "2026-03-10T07:00:00",
    "source": "STRAVA",
    "_note": "STRAVA activities are not available via the API",
}


# --- HTTP fakes -----------------------------------------------------------------------------


@dataclass
class FakeResponse:
    status_code: int = 200
    body: object = None
    headers: dict[str, str] = field(default_factory=dict)
    raw: bytes | None = None

    @property
    def content(self) -> bytes:
        return self.raw if self.raw is not None else json.dumps(self.body).encode()

    @property
    def text(self) -> str:
        return self.content.decode(errors="replace")

    def json(self) -> object:
        return self.body


@dataclass(frozen=True)
class Call:
    method: str
    url: str
    headers: Mapping[str, str]
    params: QueryParams | None
    json: object


class FakeSend:
    """Records every call; answers from a queue of responses or a handler function."""

    def __init__(
        self,
        *responses: FakeResponse,
        handler: Callable[[Call], FakeResponse] | None = None,
    ) -> None:
        self._queue = list(responses)
        self._handler = handler
        self.calls: list[Call] = []

    def __call__(
        self,
        method: str,
        url: str,
        *,
        headers: Mapping[str, str],
        params: QueryParams | None = None,
        json: object | None = None,
        timeout: float = DEFAULT_TIMEOUT_S,
    ) -> HttpResponse:
        call = Call(method, url, dict(headers), params, json)
        self.calls.append(call)
        if self._handler is not None:
            return self._handler(call)
        return self._queue.pop(0)


# --- In-memory PostgREST ----------------------------------------------------------------------

PRIMARY_KEYS: dict[str, tuple[str, ...]] = {
    "activities": ("id",),
    "strength_sets": ("sheet_id", "block", "sheet_row", "week", "set_no"),
    "blocks": ("sheet_id", "name"),
    "daily_load": ("date",),
    "planned_sessions": ("date", "name"),
    "planned_targets": ("date", "name"),
}


def _condition(row: JsonRow, column: str, op: str, value: str) -> bool:
    cell = str(row[column])
    match op:
        case "eq":
            return cell == value
        case "gte":
            return cell >= value
        case "gt":
            return cell > value
        case "lt":
            return cell < value
        case "in":
            return cell in value.strip("()").split(",")
    raise AssertionError(f"unsupported filter op {op}")


def _matches(row: JsonRow, filters: Filters) -> bool:
    for column, expr in filters:
        if column == "or":
            parts = [p.split(".", 2) for p in expr.strip("()").split(",")]
            if not any(_condition(row, c, op, v) for c, op, v in parts):
                return False
        else:
            op, _, value = expr.partition(".")
            if not _condition(row, column, op, value):
                return False
    return True


class InMemoryPostgrest(Postgrest):
    """Enough PostgREST semantics for the table modules: eq/gte/gt/lt/in/or filters,
    primary-key conflicts on insert (like HTTP 409) and merge on upsert."""

    def __init__(self) -> None:
        super().__init__("https://fake.supabase.co", "sb_secret_fake", FakeSend())
        self.tables: dict[str, list[JsonRow]] = {name: [] for name in PRIMARY_KEYS}

    def _pk(self, table: str, row: Mapping[str, object]) -> tuple[object, ...]:
        return tuple(row[c] for c in PRIMARY_KEYS[table])

    def select(
        self, table: str, *, columns: str, order: str, filters: Filters = ()
    ) -> list[JsonRow]:
        wanted = columns.split(",")
        rows = [r for r in self.tables[table] if _matches(r, filters)]
        rows.sort(key=lambda r: tuple(str(r[c]) for c in order.split(",")))
        return [{c: r[c] for c in wanted} for r in rows]

    @staticmethod
    def _same_keys(rows: Sequence[Mapping[str, object]]) -> None:
        """PostgREST bulk writes require identical keys on every row."""
        assert len({tuple(sorted(r)) for r in rows}) <= 1, "bulk rows must share keys"

    def insert(self, table: str, rows: Sequence[Mapping[str, object]]) -> None:
        self._same_keys(rows)
        existing = {self._pk(table, r) for r in self.tables[table]}
        for row in rows:
            key = self._pk(table, row)
            if key in existing:
                raise AssertionError(f"409 duplicate key {key} in {table}")
            existing.add(key)
            self.tables[table].append(dict(row))

    def upsert(self, table: str, rows: Sequence[Mapping[str, object]], *, on_conflict: str) -> None:
        assert tuple(on_conflict.split(",")) == PRIMARY_KEYS[table]
        self._same_keys(rows)
        by_key = {self._pk(table, r): r for r in self.tables[table]}
        for row in rows:
            by_key[self._pk(table, row)] = dict(row)
        self.tables[table] = list(by_key.values())

    def delete(self, table: str, filters: Filters) -> None:
        assert filters, "unfiltered delete"
        self.tables[table] = [r for r in self.tables[table] if not _matches(r, filters)]


@pytest.fixture
def db() -> InMemoryPostgrest:
    return InMemoryPostgrest()
