"""
Læser trænerens program-tab (Powerlifting Now-template) fra xlsx-eksport
og producerer en sæt-tabel med rå score pr. sæt (formlen er uændret fra strength_collector.py).

Brug:  parse_all(xlsx_bytes, scoring, issues)   (læser alle "Program - blok *"-tabs)
Celler der ikke kan læses bliver aldrig stille til 0: de tælles i ``issues`` (årsag -> antal),
og manglende/ulæselig kg gemmes som None. Selve score-formlen er uændret.
STRENGTH_K (tidl. K) ganges på i domain.strength. Muskelfaktor og kropsvægt pr. øvelse kommer fra
public.exercises via ``scoring(type, navn)`` (domain.exercises, 2026-10-12): kropsvægt-kg =
BODYWEIGHT x bodyweight_fraction; resten af formlen er uændret.
"""
import io, logging, math, re, datetime as dt
from collections import Counter
from typing import TypedDict
import openpyxl

from training_load.domain.exercises import variant_key

log = logging.getLogger(__name__)

class ParsedSet(TypedDict):
    date: dt.date
    block: str
    row: int              # 1-baseret række i fanen
    week: int             # 1-baseret uge i fanen
    section: int          # 1-baseret dag-sektion (dato-række) i fanen; kun rækkefølge, aldrig ugedag
    type: str
    name: str
    set: int
    reps: float
    logged_kg: float | None  # kg som logget, før kropsvægt; None = tom, ulæselig, eller 0 på en vægtøvelse (fx -10%-formlen før topsættet er logget). 0 er en rigtig værdi KUN på kropsvægtøvelser (kun kropsvægt)
    kg: float             # effektiv kg brugt i scoren (logged_kg + BODYWEIGHT x bodyweight_fraction for kropsvægtøvelser)
    bodyweight: bool      # øvelsen har en bodyweight_fraction i public.exercises
    muscle_factor: float  # øvelsens muskelfaktor fra public.exercises (ukendt navn: 0.6)
    sides: int            # 2: reps står pr. side og scoren tæller begge (exercises.per_side), ellers 1
    kg_factor: int        # 2: logget kg er pr. håndvægt, to holdes (exercises.kg_per_dumbbell), ellers 1
    rpe: float | None
    prescribed: str | None  # foreskrevet load-celle som tekst (fx 'RPE 7 - 8'), kun til visning
    sets_text: str | None   # sets-cellen som skrevet (fx '2'), kun til visning
    reps_text: str | None   # reps-cellen som skrevet (fx '8 - 12'), kun til visning
    e1rm: float | None      # fanens 1RM for sættets løft (SQUAT/BENCH/DEADLIFT), ellers None
    raw: dict[str, object]  # ugens celler som skrevet: sets, reps, load, weight, lsrpe, notes, mean_weight
    logged_rpe: float | None  # LSRPE (loggede RPE), kun hele/halve 1-10; ellers None (talt)
    variant: str | None       # normaliseret variantnavn fra "V:" i NOTES (videreført i blokken); None = ukendt
    variant_display: str | None  # varianten som skrevet første gang (mellemrum samlet); None = ukendt
    score: float          # rå score, før STRENGTH_K

ABS_SET_SCORE = 10*20*(0.6**2)   # fast score pr. abs-sæt (10 reps @ 20 kg RPE 6)

# RPE-ækvivalent fra %E1RM og reps (Tuchscherer, forenklet)
# e1rm% ≈ 1 / (1 + 0.0333*(reps + (10-rpe)))   -> løst for rpe
def rpe_from_pct(pct, reps):
    if not pct or pct <= 0: return None
    rir = (1/pct - 1)/0.0333 - reps
    return max(4.0, min(10.0, 10 - rir))

def mid(txt):
    """'8 - 12' -> 10, 'RPE 6 - 7' -> 6.5, 'RPE 7 - 5' (typo) -> 6, 5 -> 5"""
    if txt is None: return None
    if isinstance(txt,(int,float)): return float(txt)
    nums = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", str(txt))]
    return sum(nums)/len(nums) if nums else None

# Decimal comma only with 1-2 decimals: "110,115" is the template's list of weights, not 110.115.
_NUMBER_TEXT = re.compile(r"^\s*-?\d+(?:\.\d+|,\d{1,2})?\s*$")

def text_number(cell: object) -> object:
    """Et tal skrevet som tekst ('137.5', '137,5', '140') -> float; alt andet uændret.
    Arket kan have KG-kolonnen formateret som tekst, så indtastet kg kommer som str."""
    if isinstance(cell, str) and _NUMBER_TEXT.match(cell):
        return float(cell.strip().replace(",", "."))
    return cell

# Samme regel som text_number: decimalkomma kun med 1-2 decimaler ("110,115" er en liste).
_DECIMAL_COMMA = re.compile(r"(\d),(\d{1,2})(?!\d)")

def decimal_dots(cell: object) -> object:
    """Decimalkomma i tekst -> punktum ('RPE 7,5' -> 'RPE 7.5'), så mid() læser 7.5 og ikke
    gennemsnittet af 7 og 5 (= 6). Alt andet uændret. Bruges på load-cellen før set_score."""
    return _DECIMAL_COMMA.sub(r"\1.\2", cell) if isinstance(cell, str) else cell

# Brugerens regel (2026-10-01): arkets RPE er kun hele eller halve tal (6, 6.5, 7), på skalaen 1-10.
_RPE_NUMBER = re.compile(r"\d+(?:\.\d+)?")

def rpe_readable(load: object) -> bool:
    """En RPE-celle ('RPE 7', 'RPE 6.5 - 7', 'RPE 6,5') er læselig, når hvert tal er et helt eller
    halvt tal i 1-10. Andet ('RPE 7,8', 'RPE 7.25', 'RPE') er ulæseligt: det tælles og bruges aldrig.
    Celler uden 'RPE' (fx '-10%' eller 0.75) er ikke RPE-celler og altid læselige her."""
    if not isinstance(load, str) or "RPE" not in load:
        return True
    nums = [float(x) for x in _RPE_NUMBER.findall(str(decimal_dots(load)))]
    return bool(nums) and all((2 * n).is_integer() and 1 <= n <= 10 for n in nums)

def cell_text(cell: object) -> str | None:
    """En celle som i arket: tekst uændret, tal uden overflødige decimaler (3.0 -> '3')."""
    if cell is None or cell == "": return None
    if isinstance(cell, (int, float)) and not isinstance(cell, bool): return f"{cell:g}"
    return str(cell).strip() or None

VARIANT_LINE = re.compile(r"^\s*(?:v|variant)\s*:\s*(.*)$", re.IGNORECASE)


def split_notes(cell: object) -> tuple[str | None, str | None, int]:
    """(variant as written, the coach's note, extra V: lines) of a NOTES cell (user, 2026-10-14):
    a line "V: <variant>" or "Variant: <variant>" (any case) names the variant; the first such line
    wins (more are counted), every other line is the coach's note. Line breaks: \r\n, \r, \n."""
    text = cell_text(cell)
    if text is None:
        return None, None, 0
    variant, extra, note = None, 0, []
    for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        m = VARIANT_LINE.match(line)
        if m is None:
            if line.strip():
                note.append(line.strip())
        elif variant is None:
            variant = m.group(1).strip()
        else:
            extra += 1
    return variant, ("\n".join(note) or None), extra


def notes_text(raw: dict[str, object]) -> str | None:
    """The coach's note of a parsed set: raw["notes"] without any "V:" line; None when blank."""
    return split_notes(raw.get("notes"))[1]

def prescribed_text(load: object) -> str | None:
    """Load-cellen som i arket: tekst uændret, procent-tal (-0.1) som '-10%', andre tal som tekst."""
    if load is None or load == "": return None
    if isinstance(load, (int, float)) and not isinstance(load, bool):
        return f"{load*100:g}%" if abs(load) < 1 else f"{load:g}"
    return str(load).strip()

# Årsager i issues (nøglerne logges som de er; aldrig celleværdier).
KG_NOT_A_NUMBER = "kg not a number"   # også negativ kg ("-5"): ulæselig, aldrig brugt
SETS_REPS_NOT_A_NUMBER = "sets/reps not a number"
WEEK_DATE_NOT_A_DATE = "week date not a date"
ROW_WITHOUT_TYPE = "prescribed row without type"
RPE_NOT_HALF = "RPE not a whole or half value"
LSRPE_NOT_HALF = "logged RPE not a whole or half value"
LSRPE_COLUMN_MISSING = "no LSRPE column in the header"
RPE_UNKNOWN = "RPE unknown, set scored at RPE 6"   # en sæt-score afhænger af UNKNOWN_RPE (kg > 0)
VARIANT_EMPTY = "variant line without a name"     # "V:" uden navn: varianten fortsætter uændret
VARIANT_TWICE = "more than one variant line in a cell"   # den første vinder
RAW_KEYS = ("sets", "reps", "load", "weight", "lsrpe", "notes", "mean_weight")  # ugens 7 første kolonner


def json_cell(cell: object) -> object:
    """A cell as JSON: numbers and text as written, dates as ISO, anything else as text."""
    if cell is None or isinstance(cell, (bool, int, str)):
        return cell
    if isinstance(cell, float):
        return cell if math.isfinite(cell) else str(cell)
    if isinstance(cell, (dt.datetime, dt.date)):
        return cell.isoformat()
    return str(cell)


def logged_rpe(cell: object, issues: Counter[str]) -> float | None:
    """LSRPE: a whole or half value 1-10 ("7", "7,5", 7.5). Blank is None; anything else is None
    and counted (never guessed)."""
    if cell is None or (isinstance(cell, str) and not cell.strip()):
        return None
    value = text_number(cell)
    if (isinstance(value, bool) or not isinstance(value, (int, float))
            or not 1 <= value <= 10 or value * 2 != round(value * 2)):
        issues[LSRPE_NOT_HALF] += 1
        return None
    return float(value)

E1RM_RANGE = (20.0, 400.0)   # kg; et 1RM udenfor er aldrig et rigtigt 1RM
E1RM_OUT_OF_RANGE = "1RM out of range"

UNKNOWN_RPE = 6.0   # formlens "RPE ukendt": ingen logget, foreskrevet eller %1RM-RPE (tælles, RPE_UNKNOWN)

def pick_rpe(typ, name, reps, kg, load, e1rm, logged_rpe=None):
    """Sættets RPE og hvor den kommer fra (fix 2026-10-13: før vandt %1RM-skønnet over trænerens RPE):
      logged      LSRPE, når brugeren har logget den
      prescribed  trænerens RPE fra load-cellen ('RPE 5 - 7' -> 6); en ulæselig RPE-celle er 'unknown'
      pct_1rm     hovedløft uden RPE i cellen (fx '75%', '-10%') med kg: skøn fra fanens 1RM-tabel,
                  som trænerens % også regner fra. Fanens 1RM er ikke testet, så den bruges ALDRIG,
                  hvor træneren har foreskrevet en RPE (1 kg for lav 1RM ~ +0.4 RPE ved 77.5 kg x 8)
      unknown     ellers: (None, 'unknown'); set_score bruger UNKNOWN_RPE"""
    if logged_rpe is not None:
        return float(logged_rpe), "logged"
    if isinstance(load, str) and "RPE" in load:
        rpe = mid(load) if rpe_readable(load) else None
        return (rpe, "prescribed") if rpe is not None else (None, "unknown")
    main = typ in e1rm and "tempo" not in name.lower()   # tempo ≈ 20-25% lettere: aldrig %1RM
    rpe = rpe_from_pct(kg/e1rm[typ], reps) if (main and kg) else None
    return (rpe, "pct_1rm") if rpe is not None else (None, "unknown")

def set_score(typ, name, reps, kg, load, bodyweight_kg, factor, e1rm, logged_rpe=None):
    """Score for ét sæt: formlen fra strength_collector.py (flyttet hertil så prognosen kan genbruge
    den). kg = logget kg (0 hvis tom). bodyweight_kg (BODYWEIGHT x fraction, None = ikke en
    kropsvægtøvelse) og factor (muskelfaktor) kommer fra public.exercises; før 2026-10-12 stod de
    her som navneliste og ben-TYPEs. RPE vælges af pick_rpe (load = cellen med decimalpunktum;
    en ulæselig RPE-celle giver RPE ukendt). Returnerer (kg brugt, rpe, score, kropsvægt?)."""
    bw = bodyweight_kg is not None
    if typ == "ABS":
        score = ABS_SET_SCORE; rpe=None
    else:
        if bw: kg += bodyweight_kg
        rpe = pick_rpe(typ, name, reps, kg, load, e1rm, logged_rpe)[0]
        if rpe is None: rpe = UNKNOWN_RPE     # ukendt (fx -10% uden kg endnu); tælles i parse_tab
        score = reps * kg * (rpe/10)**2 * factor
    return kg, rpe, score, bw

def read_e1rm(rows, issues: Counter[str]) -> dict[str, float]:
    """1RM pr. løft fra fanens top (de første 15 rækker): løftets navn med tallet 5 kolonner til højre.
    Kun FØRSTE match tæller: skabelonen har længere nede en BLOCK SBD METRICS-tabel med de samme
    navne og tonnage i samme kolonne. Et tal udenfor E1RM_RANGE tælles i issues og bruges ikke."""
    e1rm = {}
    for r in rows[:15]:
        for i,c in enumerate(r):
            if c in ("SQUAT","BENCH","DEADLIFT") and c not in e1rm and i+5 < len(r) and isinstance(r[i+5],(int,float)):
                e1rm[c] = float(r[i+5])
    for lift, kg in list(e1rm.items()):
        if not E1RM_RANGE[0] <= kg <= E1RM_RANGE[1]:
            issues[E1RM_OUT_OF_RANGE] += 1
            del e1rm[lift]
    return e1rm

def parse_all(data: bytes, scoring, issues: Counter[str] | None = None,
              counts_issues=None) -> list[ParsedSet]:
    """``scoring(type, name)`` -> domain.exercises.Scoring (muscle factor, bodyweight kg) per row.
    ``counts_issues(tab)`` False: the tab is parsed but its unreadable cells aren't counted
    (the caller drops it anyway, e.g. a template tab without a block number), and a layout error
    in it (SheetLayoutError) stops nothing."""
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    sets = []
    for tab in wb.sheetnames:
        if tab.startswith("Program - blok"):
            counted = counts_issues is None or counts_issues(tab)
            try:
                sets += parse_tab(wb[tab], tab, scoring, issues if counted else Counter())
            except SheetLayoutError as e:
                if counted: raise
                # en fane uden bloknummer smides alligevel væk af kalderen: dens layout stopper intet
                log.warning("skipping tab %r without a block number: %s", tab, e)
    return sets

def row_variants(r, name, notes_cols, issues: Counter[str]):
    """(the variant in force per week, whether the week's own NOTES named one) for one sheet row
    (user, 2026-10-14): "V: <variant>" in a week's NOTES holds from that week to the end of the
    block or the next "V:"; weeks before the first are None (ukendt). "V: <the row's own name>"
    ends the variant. An empty "V:" keeps the current one (counted). ``notes_cols``: each week's
    NOTES column from the header row."""
    own = variant_key(name)
    current, states, named = None, [], []
    for col in notes_cols:
        cell = r[col] if col < len(r) else None   # ugens NOTES
        written, _note, extra = split_notes(cell)
        if extra: issues[VARIANT_TWICE] += 1
        if written is not None and not written.strip():
            issues[VARIANT_EMPTY] += 1; written = None
        if written is not None:
            key = variant_key(written)
            current = None if key == own else (key, " ".join(written.split()))
        states.append(current)
        named.append(written is not None)
    return states, named

class SheetLayoutError(ValueError):
    """The tab can't be read safely (user, 2026-10-15), so collect-strength stops instead of
    guessing: a header row not directly under a WEEK/date row, a week without one of its columns,
    a training day whose end the sheet doesn't mark, or an exercise row outside a training day
    (e.g. the coach's exercise library pasted below the last day)."""

# Overskrift -> nøgle (RAW_KEYS). En uges kolonner findes ud fra overskriftsrækken, ikke faste
# positioner (user, 2026-10-15). SETS, REPS, LOAD, WEIGHT og NOTES skal findes; LSRPE og Mean Weight
# må mangle (LSRPE tælles som LSRPE_COLUMN_MISSING, som før).
HEADER_KEYS = {"SETS": "sets", "REPS": "reps", "LOAD": "load", "WEIGHT": "weight", "LSRPE": "lsrpe",
               "NOTES": "notes", "MEAN WEIGHT": "mean_weight"}
REQUIRED_KEYS = ("sets", "reps", "load", "weight", "notes")

def is_week_row(r) -> bool:
    """Dag-sektionens WEEK/dato-række: 'WEEK 1' i rækken. Dagens dato står i kolonne B; mangler den
    (en blok, træneren har lagt ind, men ikke dateret endnu), læses dagen uden sæt, som før."""
    return bool(r) and "WEEK 1" in [str(x) for x in r]

def is_header_row(r) -> bool:
    return bool(r) and any(isinstance(c, str) and c.strip().upper() == "SETS" for c in r)

def ends_day(r) -> bool:
    """En træningsdag slutter ved næste dags 'Micro Length'- eller 'DAY'-række, eller ved
    skabelonens 'MICROCYCLE ... METRICS'-tabeller under sidste dag."""
    label = str(r[1]).strip().upper() if r and len(r) > 1 and isinstance(r[1], str) else ""
    return label in ("MICRO LENGTH", "DAY") or (not looks_like_exercise(r) and any(
        isinstance(c, str) and c.strip().upper().startswith("MICROCYCLE") for c in r))  # aldrig en NOTES-celle

def looks_like_exercise(r) -> bool:
    """En række med en øvelse: TYPE og NAME som tekst (eller en Abs-række uden TYPE)."""
    if not r or len(r) < 3 or not isinstance(r[2], str) or not r[2].strip():
        return False
    typ = str(r[1] or "").strip()
    return bool(typ) or r[2].strip().lower().startswith("abs")

def week_columns(tab, ri, header, starts):
    """Hver uges kolonne pr. overskrift (user, 2026-10-15): ugens gruppe går fra dens 'WEEK n'-
    kolonne til næste uges. Mangler SETS, REPS, LOAD, WEIGHT eller NOTES, stopper læsningen."""
    weeks = []
    for w, start in enumerate(starts):
        end = starts[w+1] if w+1 < len(starts) else len(header)
        cols = {}
        for i in range(start, min(end, len(header))):
            key = HEADER_KEYS.get(" ".join(str(header[i] or "").split()).upper())
            if key is not None and key not in cols: cols[key] = i
        missing = [k.upper() for k in REQUIRED_KEYS if k not in cols]
        if missing:
            raise SheetLayoutError(f"{tab}, row {ri}: week {w+1} has no {', '.join(missing)} "
                                   f"column under its WEEK {w+1} label")
        weeks.append(cols)
    return weeks

def parse_tab(ws, tab, scoring, issues: Counter[str] | None = None):
    issues = Counter() if issues is None else issues
    rows = list(ws.iter_rows(values_only=True))
    e1rm = read_e1rm(rows, issues)
    sets = []
    week_dates: list[object] = []   # dagens dato pr. uge (fra WEEK-rækken)
    dated = False     # dagens WEEK-række har en dato i kolonne B
    section = 0
    lsrpe_ok = True   # LSRPE læses, indtil en overskrift mangler kolonnen i en uge (talt én gang); så ikke resten af fanen
    run_head = None   # (sektion, øvelse, variant pr. uge) for den række, der startede en øvelses-run
    # En træningsdag (user, 2026-10-15): overskriftsrækken lige under en WEEK/dato-række, til næste
    # Micro Length-, DAY- eller MICROCYCLE-række. Kan grænsen ikke afgøres, stopper læsningen.
    week_row = None   # rækken med WEEK/datoer, der venter på sin overskrift
    day_start = None  # overskriftsrækken for den dag, der læses nu (None: ingen dag åben)
    started = False   # første dag er begyndt: øvelsesrækker udenfor en dag er en fejl herfra
    cols = []         # pr. uge: overskrift -> kolonne
    for ri, r in enumerate(rows, start=1):
        if is_week_row(r):
            if day_start is not None:
                raise SheetLayoutError(f"{tab}, row {ri}: a new WEEK row, but the training day from "
                                       f"row {day_start} has no end (Micro Length, DAY or MICROCYCLE row)")
            if week_row is not None:
                raise SheetLayoutError(f"{tab}, row {week_row}: no header row directly under the WEEK row")
            # dato-række: dato står 2 kolonner efter 'WEEK n'
            starts = [i for i,c in enumerate(r) if isinstance(c,str) and c.startswith("WEEK")]
            week_dates = [r[i+2] if i+2 < len(r) else None for i in starts]
            dated = len(r) > 1 and isinstance(r[1], dt.datetime)   # udateret dag: ingen sæt endnu
            week_row = ri
            section += 1
            continue
        if week_row is not None:
            if not is_header_row(r):
                raise SheetLayoutError(f"{tab}, row {week_row}: no header row (TYPE, NAME, SETS ...) "
                                       f"directly under the WEEK row")
            cols = week_columns(tab, ri, r, starts)
            if lsrpe_ok and any("lsrpe" not in c for c in cols):
                lsrpe_ok = False; issues[LSRPE_COLUMN_MISSING] += 1
            week_row, day_start, started = None, ri, True
            continue
        if is_header_row(r):
            raise SheetLayoutError(f"{tab}, row {ri}: a header row that isn't directly under a WEEK row")
        if day_start is not None and ends_day(r):
            day_start = None
            continue
        if day_start is None:
            if started and looks_like_exercise(r):
                raise SheetLayoutError(f"{tab}, row {ri}: an exercise row ({str(r[1] or '').strip()!r}, "
                                       f"{str(r[2]).strip()!r}) outside a training day")
            continue
        if not dated or not r or not r[2]: continue
        typ, name = str(r[1] or "").strip(), str(r[2]).strip()
        if typ in ("TYPE","Micro Length","DAY","LIFT","MUSCLE GROUP") or typ.startswith("MICRO") or name=="NAME": continue
        if not typ and name.lower().startswith("abs"): typ = "ABS"
        if not typ:
            if any(c["sets"] < len(r) and r[c["sets"]] not in (None,0,"") for c in cols):
                issues[ROW_WITHOUT_TYPE] += 1   # foreskrevet, men uden TYPE: ville ellers forsvinde stille
            continue
        states = row_variants(r, name, [c["notes"] for c in cols], issues)
        if run_head is not None and run_head[0] == (section, variant_key(name)):
            # En række i samme øvelses-run uden egen "V:" følger run'ens første række uge for uge.
            first_own = next((w for w, st in enumerate(states[1]) if st), len(week_dates))
            states = ([run_head[1][w] if w < first_own else states[0][w] for w in range(len(week_dates))], states[1])
        else:
            run_head = ((section, variant_key(name)), states[0])
        for w,date in enumerate(week_dates):
            variant = states[0][w]
            exercise = scoring(typ, name) if variant is None else scoring(typ, name, variant[0])
            c = cols[w]             # uge w's kolonner fra overskriften
            def at(i): return r[i] if i < len(r) else None
            nsets, reps, load, kg = at(c["sets"]), at(c["reps"]), at(c["load"]), at(c["weight"])
            # raw: cellen under hver overskrift; mangler LSRPE/Mean Weight, cellen på skabelonens
            # plads (SETS + 4 / + 6), som skrevet
            cells = [at(c[k]) if k in c else at(c["sets"] + i) for i, k in enumerate(RAW_KEYS)]
            kg = text_number(kg)
            sets_text, reps_text = cell_text(nsets), cell_text(reps)
            if nsets in (None,0,"") or reps in (None,"") : continue
            if mid(nsets) is None or mid(reps) is None:
                issues[SETS_REPS_NOT_A_NUMBER] += 1; continue
            if not isinstance(date, dt.datetime):
                issues[WEEK_DATE_NOT_A_DATE] += 1; continue
            kg_ok = isinstance(kg,(int,float)) and not isinstance(kg,bool) and kg >= 0   # -5 er ingen kg
            if not kg_ok and kg not in (None,""): issues[KG_NOT_A_NUMBER] += 1
            nsets, reps = int(mid(nsets)), mid(reps)
            kg = float(kg) if kg_ok and isinstance(kg,(int,float)) else 0.0
            logged = kg
            if not rpe_readable(load): issues[RPE_NOT_HALF] += 1   # tælles; cellens RPE bruges aldrig ("RPE ukendt")
            lsrpe = logged_rpe(cells[4], issues) if lsrpe_ok else None
            # Brugerens input pr. øvelse (2026-10-14): reps pr. side tælles for begge sider, kg pr.
            # håndvægt for begge hænder. Formlen er uændret; kun dens input ganges op. Aldrig på
            # hovedløftenes TYPEs: deres %1RM-RPE må ikke regnes på dobbelte reps eller kg.
            sides, kg_factor = (1, 1) if typ in ("SQUAT", "BENCH", "DEADLIFT") else (exercise.sides, exercise.kg_factor)
            kg, rpe, score, bw = set_score(typ, name, reps * sides, kg * kg_factor, decimal_dots(load),
                                           exercise.bodyweight_kg, exercise.muscle_factor, e1rm, lsrpe)
            if typ != "ABS" and kg and pick_rpe(typ, name, reps, kg, decimal_dots(load), e1rm, lsrpe)[0] is None:
                issues[RPE_UNKNOWN] += 1   # scoren hviler på UNKNOWN_RPE: logges, aldrig stille
            raw = {key: json_cell(cell) for key, cell in zip(RAW_KEYS, cells)}
            for s in range(nsets):
                sets.append(ParsedSet(date=date.date(), block=tab, row=ri, week=w+1, section=section, type=typ, name=name, set=s+1,
                                      reps=reps, logged_kg=logged if kg_ok and (logged or bw) else None, kg=kg, bodyweight=bw, rpe=rpe, score=round(score,1),
                                      muscle_factor=exercise.muscle_factor, sides=sides, kg_factor=kg_factor,
                                      prescribed=prescribed_text(load),
                                      sets_text=sets_text, reps_text=reps_text,
                                      e1rm=e1rm.get(typ), raw=raw, logged_rpe=lsrpe,
                                      variant=variant[0] if variant else None,
                                      variant_display=variant[1] if variant else None))
    if day_start is not None:
        raise SheetLayoutError(f"{tab}: the training day from row {day_start} has no end (no Micro "
                               f"Length, DAY or MICROCYCLE row after it)")
    if week_row is not None:
        raise SheetLayoutError(f"{tab}, row {week_row}: no header row directly under the WEEK row")
    return sets
